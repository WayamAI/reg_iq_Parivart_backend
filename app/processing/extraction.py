import io
from typing import Optional
from app.storage.base import LocalStorage

# Initialize storage
storage = LocalStorage(base_path="./storage")

async def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)
    return "\n\n".join(text_parts)

async def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file."""
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)
    return "\n\n".join(text_parts)

async def extract_text_from_html(file_bytes: bytes) -> str:
    """Extract text from an HTML file."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(file_bytes, "html.parser")
    # Remove script and style elements
    for script in soup(["script", "style"]):
        script.decompose()
    # Get text
    text = soup.get_text()
    # Clean up whitespace
    lines = (line.strip() for line in text.splitlines())
    chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
    text = "\n".join(chunk for chunk in chunks if chunk)
    return text

async def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from a plain text file."""
    return file_bytes.decode("utf-8", errors="ignore")

async def extract_text_from_file(storage_key: str, mime_type: str) -> Optional[str]:
    """
    Extract text from a file in storage based on its MIME type.
    Returns None if extraction fails or MIME type is unsupported.
    """
    try:
        file_obj = await storage.download_file(storage_key)
        file_bytes = file_obj.read()
        file_obj.close()

        if mime_type == "application/pdf":
            return await extract_text_from_pdf(file_bytes)
        elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return await extract_text_from_docx(file_bytes)
        elif mime_type == "text/html":
            return await extract_text_from_html(file_bytes)
        elif mime_type == "text/plain":
            return await extract_text_from_txt(file_bytes)
        else:
            return None
    except Exception as e:
        # Log the error and return None
        import structlog
        logger = structlog.get_logger()
        logger.error("text_extraction_failed", storage_key=storage_key, mime_type=mime_type, error=str(e))
        return None